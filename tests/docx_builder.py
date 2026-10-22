"""Builds roster .docx files shaped like the real ones: header paragraphs, then a 3-column table."""
from pathlib import Path

import docx

HEADER = [
    "СПИСОК",
    "СТУДЕНТІВ {course} КУРСУ ОСВІТНЬОГО РІВНЯ «{level}»",
    "{form} ФОРМИ ЗДОБУТТЯ ОСВІТИ",
    "ІНЖЕНЕРНОГО НАВЧАЛЬНО-НАУКОВОГО ІНСТИТУТУ ім. Ю.М.Потебні",
    "СПЕЦІАЛЬНІСТЬ: {specialty}",
    "ОСВІТНЯ ПРОГРАМА: {program}",
    "ГРУПА: {group}",
    "Термін навчання: {term}",
]


def section(*, group="6.G3.25-з", specialty="G3 ЕЛЕКТРИЧНА ІНЖЕНЕРІЯ",
            program="ЕЛЕКТРОЕНЕРГЕТИКА, ЕЛЕКТРОТЕХНІКА ТА ЕЛЕКТРОМЕХАНІКА",
            course=2, level="БАКАЛАВР", form="ЗАОЧНОЇ", term="01.09.2025-30.06.2029",
            rows=(("1", "Бадя Віталій Олегович", "Б"),), header=None):
    lines = header if header is not None else [
        line.format(course=course, level=level, form=form, specialty=specialty,
                    program=program, group=group, term=term)
        for line in HEADER
    ]
    return lines, list(rows)


def build(path: Path, *sections) -> Path:
    d = docx.Document()
    for lines, rows in sections:
        for line in lines:
            d.add_paragraph(line)
        t = d.add_table(rows=0, cols=3)
        for r in rows:
            cells = t.add_row().cells
            for c, v in zip(cells, r):
                c.text = v
        d.add_paragraph("В.о. директора ІННІ ім. Ю.М.Потебні\t\tСергій ІЛЬЇН")
    d.save(path)
    return path
