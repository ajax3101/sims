import io
from datetime import date
from decimal import Decimal

import docx
import openpyxl
import pytest
from pypdf import PdfReader

from sims.modules.reports.document import Column, Document, KeyValueBlock, TableBlock
from sims.modules.reports.exporters import to_docx, to_pdf, to_xlsx


def _doc() -> Document:
    return Document(
        title="Фінансовий звіт", meta=["Група: 6.G3.25-з"],
        blocks=[TableBlock(
            columns=[Column("name", "ПІБ"), Column("amount", "Борг, грн", kind="money")],
            rows=[{"name": "Бадя Віталій", "amount": Decimal("10.50")},
                  {"name": "Бренер Віталій", "amount": Decimal("12345.00")}],
            totals=["amount"],
        )],
    )


def test_xlsx_layout_types_and_sum_formula():
    ws = openpyxl.load_workbook(io.BytesIO(to_xlsx(_doc()))).active
    assert (ws["A1"].value, ws["A2"].value, ws["A3"].value) == ("Фінансовий звіт", "Група: 6.G3.25-з", None)
    assert [ws["A4"].value, ws["B4"].value] == ["ПІБ", "Борг, грн"]
    assert (ws["A5"].value, ws["B5"].value, ws["B6"].value) == ("Бадя Віталій", 10.5, 12345)
    assert (ws["A7"].value, ws["B7"].value) == ("Разом", "=SUM(B5:B6)")
    assert ws["B5"].number_format == "#,##0.00"
    assert (ws["A4"].font.name, ws["A4"].font.bold, ws.freeze_panes, ws.auto_filter.ref) == ("Arial", True, "A5", "A4:B6")


def test_docx_has_title_meta_and_table_with_computed_total():
    d = docx.Document(io.BytesIO(to_docx(_doc())))
    texts = [p.text for p in d.paragraphs if p.text]
    assert texts[:2] == ["Фінансовий звіт", "Група: 6.G3.25-з"]
    [table] = d.tables
    assert [[c.text for c in r.cells] for r in table.rows] == [
        ["ПІБ", "Борг, грн"], ["Бадя Віталій", "10,50"], ["Бренер Віталій", "12 345,00"], ["Разом", "12 355,50"]]


def test_pdf_renders_cyrillic_with_embedded_font():
    reader = PdfReader(io.BytesIO(to_pdf(_doc())))
    text = "".join(p.extract_text() for p in reader.pages)
    assert "Фінансовий звіт" in text and "Бренер Віталій" in text and "12 355,50" in text
    fonts = {f["/BaseFont"] for p in reader.pages for f in p["/Resources"]["/Font"].values()}
    assert any("DejaVu" in str(f) for f in fonts)


def test_key_value_block_in_every_format():
    doc = Document(title="Картка", blocks=[KeyValueBlock([("ПІБ", "Зінов’єв Роман"), ("Статус", "навчається")])])
    ws = openpyxl.load_workbook(io.BytesIO(to_xlsx(doc))).active
    assert (ws["A3"].value, ws["B3"].value) == ("ПІБ", "Зінов’єв Роман")
    d = docx.Document(io.BytesIO(to_docx(doc)))
    assert [[c.text for c in r.cells] for r in d.tables[0].rows][0] == ["ПІБ", "Зінов’єв Роман"]
    assert "Зінов’єв Роман" in PdfReader(io.BytesIO(to_pdf(doc))).pages[0].extract_text()


def test_wide_tables_switch_docx_to_landscape():
    cols = [Column(f"c{i}", f"Колонка {i}") for i in range(8)]
    d = docx.Document(io.BytesIO(to_docx(Document(title="T", blocks=[TableBlock(columns=cols, rows=[])]))))
    s = d.sections[0]
    assert s.page_width > s.page_height


@pytest.mark.parametrize("value, want", [(Decimal("0"), "0,00"), (Decimal("1234567.5"), "1 234 567,50"), (None, "")])
def test_money_text_format(value, want):
    from sims.modules.reports.document import money_text
    assert money_text(value) == want


def test_xlsx_sheet_title_drops_characters_excel_forbids():
    doc = Document(title="Звіт: борги [2026/27]? *повний* — понад тридцять один символ", blocks=[])
    ws = openpyxl.load_workbook(io.BytesIO(to_xlsx(doc))).active
    assert ws.title == "Звіт борги 202627 повний — пона"  # 31 chars, the Excel limit


def test_pdf_columns_fit_their_longest_unbreakable_word():
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfbase.pdfmetrics import stringWidth

    from sims.modules.reports.exporters import CELL_PADDING, _pdf_col_widths, _pdf_fonts

    regular, bold = _pdf_fonts()
    block = TableBlock(
        columns=[Column("n", "№", "int"), Column("name", "ПІБ"), Column("course", "Курс", "int"),
                 Column("discipline", "Дисципліна"), Column("year", "Навч. рік"), Column("semester", "Семестр", "int"),
                 Column("status", "Статус"), Column("opened", "Відкрито", "date"), Column("deadline", "Перескласти до", "date")],
        rows=[{"n": 1, "name": "Волохіна-Черняк Анна Миколаївна", "course": 3,
               "discipline": "Іноземна мова (за професійним спрямуванням)", "year": "2025/2026", "semester": 1,
               "status": "призначено перескладання", "opened": date(2026, 1, 25), "deadline": date(2026, 2, 15)}],
    )
    width = landscape(A4)[0] - 24 * 72 / 25.4
    widths = _pdf_col_widths(block, width)
    need = {  # longest token of each column, header in bold
        "Курс": stringWidth("Курс", bold, 8), "Семестр": stringWidth("Семестр", bold, 8),
        "Відкрито": stringWidth("25.01.2026", regular, 8), "Навч. рік": stringWidth("2025/2026", regular, 8),
        "Перескласти до": stringWidth("Перескласти", bold, 8), "ПІБ": stringWidth("Волохіна-Черняк", regular, 8),
    }
    got = {c.header: w for c, w in zip(block.columns, widths)}
    assert all(got[h] >= n + CELL_PADDING for h, n in need.items()), {h: (got[h], n + CELL_PADDING) for h, n in need.items()}
    assert sum(widths) <= width + 0.01
