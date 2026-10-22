import io
from datetime import date
from decimal import Decimal
from urllib.parse import unquote

import docx
import openpyxl
import pytest
from pypdf import PdfReader

from sims.models import (
    AcademicDebt, Contract, DebtStatus, Discipline, Grade, GradeKind, OrgUnit, OrgUnitKind, Payment,
    PaymentScheduleItem,
)
from tests import factories as f

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def _world(session):
    await f.academic_year(session, 2026, is_current=True)
    faculty = OrgUnit(kind=OrgUnitKind.FACULTY, name="Факультет")
    session.add(faculty)
    await session.flush()
    g = await f.group(session, code="6.G3.25-з-r1", admission_year=2025)
    (await session.get(OrgUnit, g.org_unit_id)).parent_id = faculty.id      # group's institute sits under the faculty
    dept = OrgUnit(kind=OrgUnitKind.DEPARTMENT, name="Кафедра ЕМ", parent_id=faculty.id)
    session.add(dept)
    await session.flush()
    other = await f.group(session, code="6.G3.25-з-r2")
    a = await f.student(session, last_name="Ачкасов", first_name="Іван", current_group_id=g.id)
    b = await f.student(session, last_name="Бадя", first_name="Віталій", current_group_id=g.id)
    c = await f.student(session, last_name="Чорний", first_name="Олег", current_group_id=other.id)
    gd = await f.group_discipline(session, g)
    (await session.get(Discipline, gd.discipline_id)).department_id = dept.id
    gd_other = await f.group_discipline(session, other)
    session.add_all([
        AcademicDebt(student_id=a.id, group_discipline_id=gd.id, opened_on=date(2026, 1, 25)),
        AcademicDebt(student_id=b.id, group_discipline_id=gd.id, opened_on=date(2026, 1, 25),
                     status=DebtStatus.CLOSED, closed_on=date(2026, 2, 1)),
        AcademicDebt(student_id=c.id, group_discipline_id=gd_other.id, opened_on=date(2026, 1, 25)),
        Grade(student_id=b.id, group_discipline_id=gd.id, kind=GradeKind.FINAL, score=88, ects="B",
              national="good", graded_on=date(2026, 1, 20)),
    ])
    for st, number, paid in [(a, "К-A", "1000"), (b, "К-B", "20000")]:
        contract = Contract(student_id=st.id, number=number, starts_on=date(2025, 9, 1), total_amount=Decimal("20000"))
        contract.schedule = [PaymentScheduleItem(due_date=date(2000, 9, 1), amount=Decimal("20000"))]
        contract.payments = [Payment(paid_on=date(2000, 9, 1), amount=Decimal(paid))]
        session.add(contract)
    await session.flush()
    return {"faculty": faculty, "dept": dept, "g": g, "a": a, "b": b, "c": c, "gd": gd}


def _xlsx_rows(content: bytes) -> list[list]:
    ws = openpyxl.load_workbook(io.BytesIO(content)).active
    return [[c.value for c in r] for r in ws.iter_rows()]


def _column(rows, header) -> list:
    i = next(n for n, r in enumerate(rows) if header in r)
    col = rows[i].index(header)
    return [r[col] for r in rows[i + 1:] if r[col] is not None]


async def test_academic_debt_report_for_group_lists_only_unresolved_debts(client, session):
    w = await _world(session)
    resp = await client.get("/api/reports/academic-debts", params={"group_id": str(w["g"].id), "format": "xlsx"})
    assert (resp.status_code, resp.headers["content-type"]) == (200, XLSX)
    rows = _xlsx_rows(resp.content)
    assert _column(rows, "ПІБ") == ["Ачкасов Іван"]
    assert _column(rows, "Курс") == [2]
    assert _column(rows, "Статус") == ["відкрита"]


@pytest.mark.parametrize("scope, want", [("org_unit_id", ["Ачкасов Іван"]), ("department_id", ["Ачкасов Іван"])])
async def test_academic_debt_report_by_faculty_subtree_and_by_department(client, session, scope, want):
    w = await _world(session)
    unit = w["faculty"] if scope == "org_unit_id" else w["dept"]
    resp = await client.get("/api/reports/academic-debts", params={scope: str(unit.id), "format": "xlsx"})
    assert _column(_xlsx_rows(resp.content), "ПІБ") == want


async def test_financial_report_lists_debtors_with_sum_formula(client, session):
    w = await _world(session)
    resp = await client.get("/api/reports/financial-debts", params={"group_id": str(w["g"].id), "format": "xlsx"})
    rows = _xlsx_rows(resp.content)
    assert (_column(rows, "№"), _column(rows, "ПІБ")) == ([1, "Разом"], ["Ачкасов Іван"])
    assert _column(rows, "Борг, грн") == [19000, "=SUM(H5:H5)"]


async def test_financial_report_docx_totals(client, session):
    w = await _world(session)
    resp = await client.get("/api/reports/financial-debts", params={"group_id": str(w["g"].id), "format": "docx"})
    table = docx.Document(io.BytesIO(resp.content)).tables[0]
    assert [c.text for c in table.rows[-1].cells][-1] == "19 000,00"


@pytest.mark.parametrize("fmt", ["pdf", "docx", "xlsx"])
async def test_student_card_in_every_format(client, session, fmt):
    w = await _world(session)
    resp = await client.get(f"/api/reports/students/{w['b'].id}/card", params={"format": fmt})
    assert resp.status_code == 200
    assert unquote(resp.headers["content-disposition"]).endswith(f"Бадя Віталій.{fmt}")
    if fmt == "pdf":
        text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(resp.content)).pages)
        assert "Бадя Віталій" in text and "6.G3.25-з-r1" in text and "88" in text


async def test_unknown_format_is_422(client):
    assert (await client.get("/api/reports/academic-debts", params={"format": "odt"})).status_code == 422
