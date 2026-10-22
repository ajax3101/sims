"""XLSX (openpyxl), DOCX (python-docx) and PDF (ReportLab) renderings of a report Document."""
import io
import re
from functools import cache
from pathlib import Path

import docx
import openpyxl
from docx.enum.section import WD_ORIENT
from docx.shared import Pt
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from sims.modules.reports.document import TOTAL_LABEL, Document, KeyValueBlock, TableBlock, cell_text

LANDSCAPE_FROM_COLUMNS = 7
CELL_PADDING = 12  # ReportLab Table default: 6 pt left + 6 pt right
CELL_FONT_SIZE = 8
MAX_COL_WIDTH = 220  # pt: long texts wrap rather than starve the other columns
FONTS_DIR = Path(__file__).resolve().parents[2] / "assets" / "fonts"

# --- XLSX -------------------------------------------------------------------------

_FONT = "Arial"
_NUMBER_FORMATS = {"money": "#,##0.00", "int": "0", "date": "DD.MM.YYYY"}


def to_xlsx(doc: Document) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = re.sub(r"\s+", " ", re.sub(r"[\[\]:*?/\\]", "", doc.title)).strip()[:31] or "Звіт"
    ws.cell(1, 1, doc.title).font = Font(name=_FONT, bold=True, size=14)
    row = 2
    for line in doc.meta:
        ws.cell(row, 1, line).font = Font(name=_FONT, italic=True)
        row += 1
    row += 1  # blank line before the first block
    widths: dict[int, int] = {}
    first_table = True
    for block in doc.blocks:
        if block.title:
            ws.cell(row, 1, block.title).font = Font(name=_FONT, bold=True, size=12)
            row += 1
        if isinstance(block, KeyValueBlock):
            for label, value in block.pairs:
                ws.cell(row, 1, label).font = Font(name=_FONT, bold=True)
                ws.cell(row, 2, cell_text(value)).font = Font(name=_FONT)
                widths[1] = max(widths.get(1, 0), len(label))
                widths[2] = max(widths.get(2, 0), len(cell_text(value)))
                row += 1
        else:
            row = _xlsx_table(ws, block, row, widths, first_table)
            first_table = False
        row += 1
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 8), 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_table(ws, block: TableBlock, row: int, widths: dict[int, int], first: bool) -> int:
    header_row = row
    fill = PatternFill("solid", fgColor="DDDDDD")
    for c, col in enumerate(block.columns, start=1):
        cell = ws.cell(row, c, col.header)
        cell.font, cell.fill = Font(name=_FONT, bold=True), fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        widths[c] = max(widths.get(c, 0), len(col.header))
    for data in block.rows:
        row += 1
        for c, col in enumerate(block.columns, start=1):
            value = data.get(col.key)
            cell = ws.cell(row, c, value if col.kind in _NUMBER_FORMATS else cell_text(value))
            cell.font = Font(name=_FONT)
            if col.kind in _NUMBER_FORMATS:
                cell.number_format = _NUMBER_FORMATS[col.kind]
            widths[c] = max(widths.get(c, 0), len(cell_text(value, col.kind)))
    last_data = row
    if first:
        ws.freeze_panes = ws.cell(header_row + 1, 1).coordinate
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(block.columns))}{last_data}"
    if block.totals:
        row += 1
        ws.cell(row, 1, TOTAL_LABEL).font = Font(name=_FONT, bold=True)
        for c, col in enumerate(block.columns, start=1):
            if col.key in block.totals:
                letter = get_column_letter(c)
                # A formula, not a Python-computed number: the sheet stays correct when rows are edited.
                value = f"=SUM({letter}{header_row + 1}:{letter}{last_data})" if block.rows else 0
                cell = ws.cell(row, c, value)
                cell.font, cell.number_format = Font(name=_FONT, bold=True), _NUMBER_FORMATS.get(col.kind, "General")
    return row


# --- DOCX -------------------------------------------------------------------------

def to_docx(doc: Document) -> bytes:
    d = docx.Document()
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(11)
    if doc.widest_table >= LANDSCAPE_FROM_COLUMNS:
        s = d.sections[0]
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = s.page_height, s.page_width
    d.add_heading(doc.title, level=1)
    for line in doc.meta:
        d.add_paragraph(line).runs[0].italic = True
    for block in doc.blocks:
        if block.title:
            d.add_heading(block.title, level=2)
        if isinstance(block, KeyValueBlock):
            t = d.add_table(rows=0, cols=2)
            t.style = "Table Grid"
            for label, value in block.pairs:
                cells = t.add_row().cells
                cells[0].text, cells[1].text = label, cell_text(value)
                cells[0].paragraphs[0].runs[0].bold = True
        else:
            _docx_table(d, block)
        d.add_paragraph()
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _docx_table(d, block: TableBlock) -> None:
    t = d.add_table(rows=1, cols=len(block.columns))
    t.style = "Table Grid"
    for cell, col in zip(t.rows[0].cells, block.columns):
        cell.text = col.header
        cell.paragraphs[0].runs[0].bold = True
    for data in block.rows:
        for cell, col in zip(t.add_row().cells, block.columns):
            cell.text = cell_text(data.get(col.key), col.kind)
    if block.totals:
        cells = t.add_row().cells
        cells[0].text = TOTAL_LABEL
        cells[0].paragraphs[0].runs[0].bold = True
        for cell, col in zip(cells, block.columns):
            if col.key in block.totals:
                cell.text = cell_text(block.total(col.key), col.kind)


# --- PDF --------------------------------------------------------------------------

@cache
def _pdf_fonts() -> tuple[str, str]:
    # Base-14 PDF fonts have no Cyrillic glyphs; embed DejaVu Sans.
    pdfmetrics.registerFont(TTFont("DejaVuSans", str(FONTS_DIR / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", str(FONTS_DIR / "DejaVuSans-Bold.ttf")))
    pdfmetrics.registerFontFamily("DejaVuSans", normal="DejaVuSans", bold="DejaVuSans-Bold")
    return "DejaVuSans", "DejaVuSans-Bold"


def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def to_pdf(doc: Document) -> bytes:
    regular, bold = _pdf_fonts()
    styles = {
        "title": ParagraphStyle("title", fontName=bold, fontSize=14, leading=18, spaceAfter=4),
        "meta": ParagraphStyle("meta", fontName=regular, fontSize=9, leading=12, textColor=colors.grey),
        "h2": ParagraphStyle("h2", fontName=bold, fontSize=11, leading=14, spaceBefore=6, spaceAfter=4),
        "cell": ParagraphStyle("cell", fontName=regular, fontSize=CELL_FONT_SIZE, leading=10),
        "head": ParagraphStyle("head", fontName=bold, fontSize=CELL_FONT_SIZE, leading=10),
    }
    pagesize = landscape(A4) if doc.widest_table >= LANDSCAPE_FROM_COLUMNS else A4
    buf = io.BytesIO()
    frame = SimpleDocTemplate(buf, pagesize=pagesize, leftMargin=12 * mm, rightMargin=12 * mm,
                              topMargin=12 * mm, bottomMargin=12 * mm, title=doc.title)
    story = [Paragraph(_esc(doc.title), styles["title"])]
    story += [Paragraph(_esc(line), styles["meta"]) for line in doc.meta]
    story.append(Spacer(1, 4 * mm))
    for block in doc.blocks:
        if block.title:
            story.append(Paragraph(_esc(block.title), styles["h2"]))
        if isinstance(block, KeyValueBlock):
            data = [[Paragraph(_esc(k), styles["head"]), Paragraph(_esc(cell_text(v)), styles["cell"])]
                    for k, v in block.pairs]
            widths = [frame.width * 0.3, frame.width * 0.7]
            story.append(_pdf_table(data, widths, header=False))
        else:
            story.append(_pdf_table(*_pdf_table_data(block, styles, frame.width), header=True))
        story.append(Spacer(1, 4 * mm))
    frame.build(story)
    return buf.getvalue()


def _pdf_table_data(block: TableBlock, styles, width: float):
    head = [Paragraph(_esc(c.header), styles["head"]) for c in block.columns]
    rows = [[Paragraph(_esc(cell_text(r.get(c.key), c.kind)), styles["cell"]) for c in block.columns]
            for r in block.rows]
    if block.totals:
        rows.append([Paragraph(TOTAL_LABEL if i == 0 else
                               (_esc(cell_text(block.total(c.key), c.kind)) if c.key in block.totals else ""),
                               styles["head"]) for i, c in enumerate(block.columns)])
    return [head, *rows], _pdf_col_widths(block, width)


def _pdf_col_widths(block: TableBlock, width: float) -> list[float]:
    """Each column at least as wide as its longest unbreakable word; spare width goes where text is long."""
    regular, bold = _pdf_fonts()

    def measure(text: str, font: str) -> float:
        return stringWidth(text, font, CELL_FONT_SIZE)

    mins, wants = [], []
    for c in block.columns:
        texts = [cell_text(r.get(c.key), c.kind) for r in block.rows]
        words = [(w, regular) for t in texts for w in t.split()] + [(w, bold) for w in c.header.split()]
        lo = max((measure(w, f) for w, f in words), default=0) + CELL_PADDING
        hi = min(max([measure(c.header, bold)] + [measure(t, regular) for t in texts]) + CELL_PADDING, MAX_COL_WIDTH)
        mins.append(lo)
        wants.append(max(hi, lo))
    if sum(mins) >= width:  # cannot honour every minimum: shrink proportionally
        return [m * width / sum(mins) for m in mins]
    if sum(wants) <= width:  # everything fits on one line: share the rest by desired width
        return [w * width / sum(wants) for w in wants]
    slack = [w - m for w, m in zip(wants, mins)]
    return [m + (width - sum(mins)) * sl / sum(slack) for m, sl in zip(mins, slack)]


def _pdf_table(data, widths, header: bool) -> Table:
    t = Table(data, colWidths=widths, repeatRows=1 if header else 0)
    style = [("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]
    if header:
        style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDDDDD")))
    t.setStyle(TableStyle(style))
    return t
